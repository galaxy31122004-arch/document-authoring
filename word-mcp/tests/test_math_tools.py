from docx import Document

from word_document_server.tools.math_tools import (
    format_math_expression,
    format_math_expression_in_table_cell,
)


def test_subscript_preserves_surrounding_formatting(tmp_path):
    path = tmp_path / "math.docx"
    doc = Document()
    paragraph = doc.add_paragraph()
    paragraph.add_run("- ")
    heading = paragraph.add_run("Phân tích: ")
    heading.bold = True
    paragraph.add_run("Tần số f_")
    paragraph.add_run("CK = 2400 Hz.")
    doc.save(path)

    result = format_math_expression(str(path), 0, "f_{CK}")
    assert result.startswith("Formatted")
    paragraph = Document(path).paragraphs[0]
    assert paragraph.text == "- Phân tích: Tần số fCK = 2400 Hz."
    assert any(run.text == "CK" and run.font.subscript for run in paragraph.runs)
    assert any(run.text == "Phân tích: " and run.bold for run in paragraph.runs)


def test_superscript_and_occurrence(tmp_path):
    path = tmp_path / "math.docx"
    doc = Document()
    doc.add_paragraph("x^2 + x^2")
    doc.save(path)

    result = format_math_expression(str(path), 0, "x^{2}", occurrence=2)
    assert result.startswith("Formatted")
    paragraph = Document(path).paragraphs[0]
    assert paragraph.text == "x^2 + x2"
    assert any(run.text == "2" and run.font.superscript for run in paragraph.runs)


def test_invalid_notation_does_not_change_document(tmp_path):
    path = tmp_path / "math.docx"
    doc = Document()
    doc.add_paragraph("f_CK")
    doc.save(path)
    before = path.read_bytes()

    assert format_math_expression(str(path), 0, "f_CK + x^2").startswith("Invalid notation")
    assert path.read_bytes() == before


def test_braced_subscript_and_table_cell(tmp_path):
    path = tmp_path / "math.docx"
    doc = Document()
    doc.add_paragraph("b_{k-1}")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "D = Delta f / f_b"
    doc.save(path)

    assert format_math_expression(str(path), 0, "b_{k-1}").startswith("Formatted")
    assert format_math_expression_in_table_cell(str(path), 0, 0, 0, "f_{b}").startswith("Formatted")
    doc = Document(path)
    assert doc.paragraphs[0].text == "bk-1"
    assert any(run.text == "k-1" and run.font.subscript for run in doc.paragraphs[0].runs)
    assert doc.tables[0].cell(0, 0).text == "D = Delta f / fb"
    assert any(run.text == "b" and run.font.subscript
               for run in doc.tables[0].cell(0, 0).paragraphs[0].runs)
