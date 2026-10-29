"""Inline mathematical notation tools for editable Word text."""

import copy
import os
import re

from docx import Document
from docx.text.run import Run

from word_document_server.utils.file_utils import check_file_writeable, ensure_docx_extension


_NOTATION = re.compile(r"^([A-Za-z][A-Za-z0-9]*)([_^])(?:\{([A-Za-z0-9+-]+)\}|([A-Za-z0-9]+))$")


def format_math_expression(filename: str, paragraph_index: int, notation: str,
                           occurrence: int = 1) -> str:
    """Format one inline subscript/superscript without changing surrounding runs.

    ``f_{CK}`` matches ``f_CK`` in the document and displays CK as a Word
    subscript. ``x^{2}`` similarly formats a superscript. The braces are input
    syntax only; the document may contain the plain or braced spelling.
    """
    if _NOTATION.fullmatch(notation) is None:
        return "Invalid notation. Use f_{CK}, b_{k-1}, or x^{2}."
    try:
        paragraph_index = int(paragraph_index)
        occurrence = int(occurrence)
    except (TypeError, ValueError):
        return "Invalid paragraph_index or occurrence: both must be integers."
    if occurrence < 1:
        return "Invalid occurrence: use 1 or greater."

    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return f"Document {filename} does not exist"
    writable, error = check_file_writeable(filename)
    if not writable:
        return f"Cannot modify document: {error}. Consider creating a copy first."

    try:
        doc = Document(filename)
        if not 0 <= paragraph_index < len(doc.paragraphs):
            return f"Invalid paragraph index. Document has {len(doc.paragraphs)} paragraphs."
        paragraph = doc.paragraphs[paragraph_index]
        result = _format_paragraph(paragraph, notation, occurrence)
        if not result.startswith("Formatted"):
            return result
        doc.save(filename)
        return f"Formatted {notation} occurrence {occurrence} in paragraph {paragraph_index}."
    except Exception as exc:
        return f"Failed to format expression: {exc}"


def format_math_expression_in_table_cell(filename: str, table_index: int,
                                         row_index: int, col_index: int, notation: str,
                                         paragraph_index: int = 0,
                                         occurrence: int = 1) -> str:
    """Apply inline script formatting to one paragraph inside a table cell."""
    if _NOTATION.fullmatch(notation) is None:
        return "Invalid notation. Use f_{CK}, b_{k-1}, or x^{2}."
    try:
        table_index, row_index, col_index, paragraph_index, occurrence = map(
            int, (table_index, row_index, col_index, paragraph_index, occurrence)
        )
    except (TypeError, ValueError):
        return "Invalid index or occurrence: use integers."
    if occurrence < 1:
        return "Invalid occurrence: use 1 or greater."

    filename = ensure_docx_extension(filename)
    if not os.path.exists(filename):
        return f"Document {filename} does not exist"
    writable, error = check_file_writeable(filename)
    if not writable:
        return f"Cannot modify document: {error}. Consider creating a copy first."
    try:
        doc = Document(filename)
        cell = doc.tables[table_index].cell(row_index, col_index)
        paragraph = cell.paragraphs[paragraph_index]
        result = _format_paragraph(paragraph, notation, occurrence)
        if not result.startswith("Formatted"):
            return result
        doc.save(filename)
        return (f"Formatted {notation} occurrence {occurrence} in table {table_index}, "
                f"row {row_index}, column {col_index}, paragraph {paragraph_index}.")
    except IndexError:
        return "Invalid table, row, column, or paragraph index."
    except Exception as exc:
        return f"Failed to format expression: {exc}"


def _format_paragraph(paragraph, notation: str, occurrence: int) -> str:
    if "".join(run.text for run in paragraph.runs) != paragraph.text:
        return "Paragraph contains unsupported text containers; no changes made."
    match = _NOTATION.fullmatch(notation)
    base, marker = match.group(1, 2)
    script = match.group(3) or match.group(4)
    candidates = (f"{base}{marker}{script}", f"{base}{marker}{{{script}}}")
    matches = sorted((m.start(), m.end(), spelling)
                     for spelling in candidates
                     for m in re.finditer(re.escape(spelling), paragraph.text))
    if len(matches) < occurrence:
        return f"Expression {notation} occurrence {occurrence} not found."
    start, end, spelling = matches[occurrence - 1]
    script_start = start + len(base) + 1 + (1 if "{" in spelling else 0)
    script_end = script_start + len(script)
    flags = [None] * len(paragraph.text)
    flags[start + len(base)] = "skip"
    if "{" in spelling:
        flags[script_start - 1] = "skip"
        flags[script_end] = "skip"
    for position in range(script_start, script_end):
        flags[position] = "subscript" if marker == "_" else "superscript"

    offset = 0
    for run in list(paragraph.runs):
        run_end = offset + len(run.text)
        if run_end <= start or offset >= end:
            offset = run_end
            continue
        if any(child.tag.rsplit("}", 1)[-1] not in ("rPr", "t") for child in run._r):
            return "Expression overlaps a non-text run; no changes made."
        groups = []
        for local, char in enumerate(run.text):
            flag = flags[offset + local]
            if flag == "skip":
                continue
            if groups and groups[-1][0] == flag:
                groups[-1] = (flag, groups[-1][1] + char)
            else:
                groups.append((flag, char))
        for flag, segment in groups:
            element = copy.deepcopy(run._r)
            fragment = Run(element, paragraph)
            fragment.text = segment
            if flag == "subscript":
                fragment.font.subscript = True
            elif flag == "superscript":
                fragment.font.superscript = True
            run._r.addprevious(element)
        run._r.getparent().remove(run._r)
        offset = run_end
    return "Formatted"
